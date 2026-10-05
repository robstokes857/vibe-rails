"""Export the v0.2 manuscript without rerunning research or rebuilding its PDF."""
from pathlib import Path
import ast, copy, hashlib, html, json, re
from lxml import etree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output/docx'; OUT.mkdir(parents=True,exist_ok=True)
src=(ROOT/'build_paper.py').read_text(encoding='utf-8')
data=json.loads((ROOT/'evidence.json').read_text(encoding='utf-8'))
env={'recent':data['september8_bounded_requests'],'u':data['september8_claude_reported_usage'],
     'hashlib':hashlib,'DOSSIER':Path('C:/source/vibe-books/token_saver/research_2026-10-01.md')}
exec(compile(src[src.index('pages=[]'):src.index('# Capture provenance')],'<manuscript blocks>','exec'),env)
pages=env['pages']
for title,blocks in pages:
    if title.startswith('13.'):
        package=next(b for b in blocks if b[0]=='table')
        package[2].extend([
            ['output/docx/tokensaver-research-editable.docx','Word review copy with native equations and tracked future edits.'],
            ['tokensaver-research-review.md / REVIEW.md','Agent-editable manuscript, stable claim IDs and return instructions.']])
(ROOT/'manuscript-blocks.json').write_text(json.dumps(pages,indent=2),encoding='utf-8')
symbols=ast.literal_eval(next(n.value for n in ast.parse(src).body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='inline_symbols' for t in n.targets)))

def mdtext(s):
    s=html.unescape(s).replace('<b>','**').replace('</b>','**')
    for key,(_,latex) in symbols.items(): s=s.replace(key,'$'+latex+'$')
    return s

md=[]
for i,(title,blocks) in enumerate(pages):
    md.append(('# ' if i==0 else '## ')+title+'\n')
    for b in blocks:
        if b[0] in ('p','note','subtitle'):md.append(mdtext(b[1])+'\n')
        elif b[0]=='h':md.append('### '+b[1]+'\n')
        elif b[0]=='eq':md.append('$$\n'+b[1]+r'\tag{'+str(b[2])+'}\n$$\n')
        elif b[0]=='fig':md.append('!['+b[2].split('.')[0]+'](figures/'+b[1]+'.png)\n\n'+b[2]+'\n')
        elif b[0]=='table':
            md.append('| '+' | '.join(b[1])+' |\n| '+' | '.join('---' for _ in b[1])+' |\n'+
                      '\n'.join('| '+' | '.join(str(v).replace('|',r'\|') for v in row)+' |' for row in b[2])+'\n')
md_path=ROOT/'tokensaver-research-review.md'
md_path.write_text('\n'.join(md),encoding='utf-8')

# Native Word math via Microsoft's installed MathML to OMML transform.
M='http://www.w3.org/1998/Math/MathML'
def node(tag,*children,text=None):
    n=ET.Element('{'+M+'}'+tag)
    if text is not None:n.text=text
    for c in children:n.append(copy.deepcopy(c))
    return n
def mi(s):return node('mi',text=s)
def mo(s):return node('mo',text=s)
def mn(s):return node('mn',text=str(s))
def row(*args):return node('mrow',*args)
def sub(s,t):return node('msub',mi(s),mi(t))
def par(x):return row(mo('('),x,mo(')'))
def summand(x,index='j',upper='J'):
    sign=node('munderover',mo('∑'),row(mi(index),mo('='),mn(1)),mi(upper)) if upper else node('munder',mo('∑'),mi(index))
    # Microsoft XSLT consumes an explicit mrow as the n-ary argument.
    return row(sign,row(x))
def frac(a,b):return node('mfrac',a,b)
def norm(x):return row(mo('|'),x,mo('|'))
space=node('mspace');space.set('width','1em')
bodyB=node('msubsup',mi('B'),mi('j'),node('mtext',text='body'))
bodyA=node('msubsup',mi('A'),mi('j'),node('mtext',text='body'))
eqs={
1:row(sub('D','B'),mo('='),summand(par(row(sub('B','j'),mo('−'),sub('A','j')))),mo(','),space,sub('S','B'),mo('='),frac(sub('D','B'),summand(sub('B','j')))),
2:row(sub('D','T'),mo('='),summand(row(mo('['),sub('τ','m'),par(bodyB),mo('−'),sub('τ','m'),par(bodyA),mo(']')))),
3:row(norm(row(mi('R'),par(mi('B')))),mo('='),summand(norm(sub('v','i')),'i',None),mo('+'),summand(norm(sub('a','i')),'i',None),mo('≤'),summand(norm(sub('v','i')),'i',None),mo('+'),summand(norm(sub('s','i')),'i',None),mo('='),norm(mi('B'))),
4:row(mi('E'),mo('['),mi('N'),mo(']'),mo('='),mi('G'),mo('−'),mi('q'),mi('C'),mo('−'),mi('K'),mo(','),space,node('msup',mi('q'),mo('*')),mo('='),frac(row(mi('G'),mo('−'),mi('K')),mi('C')),space,par(row(mi('C'),mo('>'),mn(0)))),
5:row(mo('∃'),mi('h'),mo(':'),space,mi('h'),par(row(mi('C'),par(mi('x')))),mo('='),mi('g'),par(mi('x')),space,mo('∀'),mi('x'),space,mo('⟺'),space,mi('C'),par(mi('x')),mo('='),mi('C'),par(mi('y')),mo('⇒'),mi('g'),par(mi('x')),mo('='),mi('g'),par(mi('y'))),
6:row(node('msub',node('mover',mi('S'),mo('^')),node('mtext',text='cost')),mo('='),mn(1),mo('−'),frac(summand(sub('C','k,C'),'k','n'),summand(sub('C','k,A'),'k','n')))
}
transform=ET.XSLT(ET.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
def omml(expr):
    math=node('math',expr);math.set('display','block')
    result=transform(math)
    return copy.deepcopy(result.getroot())

doc=Document(); sec=doc.sections[0]
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',25),('Subtitle',13),('Heading 1',17),('Heading 2',12)]:
    s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else '17324D')
    s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.keep_with_next=True
for name in ['Caption']:
    doc.styles[name].font.name='Calibri';doc.styles[name].font.size=Pt(9)
    doc.styles[name].font.color.rgb=RGBColor.from_string('526779')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('TokenSaver research 0.2  |  ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Measuring TokenSaver Compression Savings'
doc.core_properties.subject='Editable research manuscript for evidence review'
doc.core_properties.author='Prepared for Robert Stokes'

def cleanheading(s):return re.sub(r'[^A-Za-z0-9 ]','',s.replace('-',' ')).strip()
def rich(p,text):
    for key,(formatted,_) in symbols.items():text=text.replace(key,formatted)
    text=text.replace('alpha','α').replace('delta','δ')
    parts=re.split(r'(<b>|</b>|<sub>|</sub>)',html.unescape(text));bold=False;subscript=False
    for t in parts:
        if t=='<b>':bold=True
        elif t=='</b>':bold=False
        elif t=='<sub>':subscript=True
        elif t=='</sub>':subscript=False
        elif t:
            r=p.add_run(t);r.bold=bold;r.font.subscript=subscript
    return p

for i,(title,blocks) in enumerate(pages):
    titlep=doc.add_paragraph(cleanheading(title),'Title' if i==0 else 'Heading 1')
    if i: titlep.paragraph_format.page_break_before=True
    for b in blocks:
        kind=b[0]
        if kind in ('p','note','subtitle'):
            p=rich(doc.add_paragraph(style='Subtitle' if kind=='subtitle' else 'Normal'),b[1])
            if kind=='note':
                for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('526779')
        elif kind=='h':doc.add_paragraph(cleanheading(b[1]),'Heading 2')
        elif kind=='eq':
            p=doc.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(10)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            math=omml(eqs[b[2]])
            assert not math.xpath('.//m:nary/m:e[not(*)]',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}), 'Empty sum operand'
            # The Microsoft transform returns a native m:oMath node.
            p._p.append(math);p.add_run('    ('+str(b[2])+')').font.size=Pt(9)
        elif kind=='fig':
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run();shape=r.add_picture(str(ROOT/'figures'/(b[1]+'.png')),width=Inches(6.2))
            shape._inline.docPr.set('descr',b[2])
            cap=doc.add_paragraph(b[2],'Caption');cap.paragraph_format.space_after=Pt(9)
        elif kind=='table':
            t=doc.add_table(rows=1,cols=len(b[1]));t.autofit=False
            widths=b[3] or [470/len(b[1])]*len(b[1]);total=sum(widths)
            for j,w in enumerate(widths):t.columns[j].width=Inches(7.1*w/total)
            for j,h in enumerate(b[1]):t.rows[0].cells[j].text=h
            head=OxmlElement('w:tblHeader');head.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(head)
            for rowdata in b[2]:
                for c,val in zip(t.add_row().cells,rowdata):c.text=str(val)
            for ir,tr in enumerate(t.rows):
                pr=tr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                for j,c in enumerate(tr.cells):
                    c.width=Inches(7.1*widths[j]/total)
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'17324D' if ir==0 else ('EDF3F7' if ir%2 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shading)
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        v=OxmlElement('w:'+side);v.set(qn('w:w'),'85');v.set(qn('w:type'),'dxa');margins.append(v)
                    c._tc.get_or_add_tcPr().append(margins)
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
                        for r in p.runs:r.font.size=Pt(9.5);r.bold=ir==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ir==0 else '17324D')
            doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Future edits are tracked; the exported baseline has no tracked insertions.
track=OxmlElement('w:trackRevisions');doc.settings.element.append(track)
docx_path=OUT/'tokensaver-research-editable.docx';doc.save(docx_path)
baseline={'paper_version':'0.2','export_date':'2026-10-05','section_count':len(pages),'native_equations':len(eqs),'figures':9,
          'markdown_sha256':hashlib.sha256(md_path.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(docx_path.read_bytes()).hexdigest(),
          'word_track_changes_enabled':True,'note':'Formatting export; no new measurement or change to scientific claims.'}
(ROOT/'review-baseline.json').write_text(json.dumps(baseline,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(docx_path),'markdown':str(md_path),'baseline':baseline},indent=2))
